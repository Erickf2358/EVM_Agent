import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from projects.models import Project

from .models import CBSControlAccount, CBSProjectGroup, WorkPackage

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def read_sheet(content):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


def sheet_to_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(row=r, column=c, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def as_upload(content):
    return SimpleUploadedFile('data.xlsx', content, XLSX)


class PrefilledTemplateTests(TestCase):
    def setUp(self):
        self.project = Project.objects.create(code='P1', name='Test Project')
        self.pg1 = CBSProjectGroup.objects.create(project=self.project, code='10', description='Structure')
        self.pg2 = CBSProjectGroup.objects.create(project=self.project, code='20', description='MEP')
        self.ca1 = CBSControlAccount.objects.create(
            project_group=self.pg1, code='100', description='Concrete'
        )
        self.ca2 = CBSControlAccount.objects.create(
            project_group=self.pg2, code='200', description='Ducts'
        )
        self.ca3 = CBSControlAccount.objects.create(
            project_group=self.pg1, code='300', description='No dates'
        )
        WorkPackage.objects.create(
            control_account=self.ca1, code='100.01', name='Pour footing',
            budget='1500.50', unit='m3', qty='30', bl_start='2026-01-05', bl_end='2026-03-20',
        )
        WorkPackage.objects.create(
            control_account=self.ca3, code='300.01', name='No baseline dates', budget='10',
        )

    def template(self, name, project=None):
        project = project or self.project
        res = self.client.get(f'/api/cbs/{name}/template/?project={project.pk}')
        self.assertEqual(res.status_code, 200)
        return read_sheet(res.content)

    def test_empty_project_gets_blank_templates(self):
        empty = Project.objects.create(code='P2', name='Empty')
        expected_widths = {'project-groups': 2, 'control-accounts': 3, 'work-packages': 8}
        for name, width in expected_widths.items():
            rows = self.template(name, project=empty)
            self.assertEqual(len(rows), 6, name)  # header + 5 usable blank rows
            for row in rows[1:]:
                self.assertEqual(row, [None] * width, name)

    def test_project_group_template_is_prefilled(self):
        rows = self.template('project-groups')
        self.assertEqual(rows[0], ['Code', 'Description'])
        self.assertEqual(rows[1:3], [['10', 'Structure'], ['20', 'MEP']])
        self.assertEqual(rows[3:], [[None, None]] * 5)

    def test_control_account_template_is_prefilled_with_pg_code(self):
        rows = self.template('control-accounts')
        self.assertEqual(rows[0], ['CBS PG', 'Code', 'Description'])
        self.assertEqual(
            rows[1:4],
            [['10', '100', 'Concrete'], ['10', '300', 'No dates'], ['20', '200', 'Ducts']],
        )
        self.assertEqual(rows[4:], [[None, None, None]] * 5)

    def test_work_package_template_is_prefilled_and_excludes_cost_activities(self):
        rows = self.template('work-packages')
        self.assertEqual(
            rows[0],
            ['CBS CA', 'CBS WP', 'WP Name', 'Budget', 'Unit', 'Qty', 'BL Start', 'BL End'],
        )
        data = [r for r in rows[1:] if r[1] is not None]
        # Qty defaults to 0 and exports as 0 rather than blank; Unit '' exports as blank.
        self.assertEqual(
            data,
            [
                ['100', '100.01', 'Pour footing', 1500.5, 'm3', 30, '2026-01-05', '2026-03-20'],
                ['300', '300.01', 'No baseline dates', 10, None, 0, None, None],
            ],
        )
        self.assertNotIn('WP-100.00', [r[1] for r in data])

    def test_template_requires_and_validates_project(self):
        self.assertEqual(self.client.get('/api/cbs/project-groups/template/').status_code, 400)
        self.assertEqual(self.client.get('/api/cbs/project-groups/template/?project=9999').status_code, 404)
        self.assertEqual(self.client.get('/api/cbs/project-groups/template/?project=abc').status_code, 404)

    def test_reexporting_work_package_template_is_a_no_op(self):
        """Re-uploading the pre-filled Work Package template must not change data."""
        fields = ['code', 'name', 'budget', 'unit', 'qty', 'bl_start', 'bl_end', 'is_cost_activity']
        before = list(WorkPackage.objects.values(*fields).order_by('id'))

        res = self.client.get(f'/api/cbs/work-packages/template/?project={self.project.pk}')

        preview = self.client.post(
            '/api/cbs/work-packages/import/preview/',
            {'project': self.project.pk, 'file': as_upload(res.content)},
        ).json()
        self.assertEqual(preview['created'], 0)
        self.assertEqual(preview['updated'], 2)
        self.assertEqual(preview['errors'], [])

        self.client.post(
            '/api/cbs/work-packages/import/',
            {'project': self.project.pk, 'file': as_upload(res.content)},
        )
        self.assertEqual(list(WorkPackage.objects.values(*fields).order_by('id')), before)
        self.assertEqual(WorkPackage.objects.filter(is_cost_activity=True).count(), 3)

    def test_typed_into_blank_row_creates_new_group(self):
        """A value typed into one of the appended blank rows must import cleanly."""
        res = self.client.get(f'/api/cbs/project-groups/template/?project={self.project.pk}')
        rows = read_sheet(res.content)
        rows[-1] = ['30', 'Interiors']

        body = self.client.post(
            '/api/cbs/project-groups/import/',
            {'project': self.project.pk, 'file': as_upload(sheet_to_bytes(rows))},
        ).json()
        self.assertEqual(body['created'], 1)
        self.assertEqual(body['updated'], 2)
        self.assertEqual(body['total'], 3)
        self.assertTrue(
            CBSProjectGroup.objects.filter(project=self.project, code='30', description='Interiors').exists()
        )